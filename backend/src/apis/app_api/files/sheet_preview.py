"""Read an .xlsx workbook into rows of display strings for the UI's grid.

This is the server-side half of the spreadsheet preview. The browser gets
no renderer and no workbook bytes — just headers and rows it draws in the
same grid the `.csv` preview uses.

**Why this runs here and not in the browser.** Every client-side option
was rejected on its own terms: the npm build of SheetJS is frozen on a
2022 release with unfixed advisories, and ExcelJS — the only base for a
maintained grid renderer — raises on any workbook containing a native
chart, which is exactly what `create_excel_spreadsheet` produces. Both of
those are *renderers*, reproducing Excel's own layout. A grid of values
needs only a reader, and openpyxl is already the library our own
spreadsheet tools drive inside Code Interpreter.

**What this deliberately is not.** No fills, fonts, borders, merges,
column widths or charts. Download-and-open remains the path for anyone
who needs fidelity. Cells arrive as strings, already formatted for
display, because the alternative is shipping a type tag per cell and
re-implementing the same formatting decisions in TypeScript.
"""

import logging
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any, Iterable, Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from apis.shared.files.models import SheetPreview

logger = logging.getLogger(__name__)


# Workbook bytes we are willing to pull into memory and parse. Well above
# anything the spreadsheet tools generate, and below the point where a
# request would threaten the container.
MAX_WORKBOOK_BYTES = 25 * 1024 * 1024

# Per-sheet caps. A docked pane shows about thirty rows at a time, so
# these are generous for previewing while keeping the JSON small.
MAX_ROWS_PER_SHEET = 500
MAX_COLUMNS_PER_SHEET = 64

# Sheets read from one workbook.
MAX_SHEETS = 12

# Global cell budget across every sheet, so a workbook of many wide
# sheets cannot multiply the per-sheet caps into a huge response.
MAX_TOTAL_CELLS = 40_000


class WorkbookTooLargeError(Exception):
    """The workbook is past `MAX_WORKBOOK_BYTES`."""

    def __init__(self, size_bytes: int):
        self.size_bytes = size_bytes
        super().__init__(
            f"Workbook is {size_bytes} bytes, over the {MAX_WORKBOOK_BYTES} limit"
        )


class WorkbookUnreadableError(Exception):
    """The bytes are not a workbook openpyxl can read."""


def read_workbook_preview(data: bytes) -> list[SheetPreview]:
    """Read every visible sheet of an .xlsx into display-ready rows.

    Raises:
        WorkbookTooLargeError: past the byte cap.
        WorkbookUnreadableError: corrupt, encrypted, or not OOXML.
    """
    if len(data) > MAX_WORKBOOK_BYTES:
        raise WorkbookTooLargeError(len(data))

    # Two passes over the same bytes, and both are needed.
    #
    # `data_only=True` yields the value Excel last cached for a formula
    # cell — and `None` when nothing was cached, which is the normal case
    # for a workbook openpyxl itself wrote, since openpyxl does no
    # evaluation. Every formula in a workbook our own tools generated
    # reads back as None, so a values-only preview would show blanks
    # exactly where the totals are.
    #
    # `data_only=False` yields the formula text for those cells. Showing
    # "=SUM(B2:B10)" is both more useful than an empty cell and more
    # honest: the file really does not carry that number yet.
    values = _load(data, data_only=True)
    formulas = _load(data, data_only=False)

    try:
        sheets: list[SheetPreview] = []
        budget = MAX_TOTAL_CELLS

        for index, value_sheet in enumerate(values.worksheets):
            if index >= MAX_SHEETS:
                break
            # Hidden sheets are hidden for a reason and Excel does not
            # show them either. A preview that surfaced scratch sheets
            # would misrepresent the workbook.
            if getattr(value_sheet, "sheet_state", "visible") != "visible":
                continue

            formula_sheet = (
                formulas.worksheets[index] if index < len(formulas.worksheets) else None
            )
            sheet, used = _read_sheet(value_sheet, formula_sheet, budget)
            budget -= used
            sheets.append(sheet)
            if budget <= 0:
                break

        return sheets
    finally:
        # `read_only` workbooks hold the zip open until closed.
        values.close()
        formulas.close()


def _load(data: bytes, *, data_only: bool):
    """Open the workbook in streaming mode, or say it is unreadable.

    `read_only=True` streams rows instead of building the whole object
    graph, which is what makes a cap-and-stop read cheap on a large
    sheet. `keep_links=False` avoids resolving external workbook
    references — we never show them, and resolving them is work.
    """
    try:
        return load_workbook(
            BytesIO(data),
            read_only=True,
            data_only=data_only,
            keep_links=False,
        )
    except Exception as e:  # openpyxl raises a wide range on bad input
        raise WorkbookUnreadableError(str(e)) from e


def _read_sheet(value_sheet, formula_sheet, budget: int) -> tuple[SheetPreview, int]:
    """Read one worksheet into a `SheetPreview`, stopping at the caps.

    Returns the preview and how much of the cell budget it consumed.
    """
    row_limit = min(MAX_ROWS_PER_SHEET, max(1, budget // MAX_COLUMNS_PER_SHEET))

    value_rows = _take_rows(value_sheet, row_limit + 1)
    formula_rows = _take_rows(formula_sheet, row_limit + 1) if formula_sheet else []

    merged = [
        _merge_row(value_rows[i], formula_rows[i] if i < len(formula_rows) else ())
        for i in range(len(value_rows))
    ]

    # Trailing empty columns are an artifact of a sheet's used range
    # reaching further than its data — a stray format on an empty cell is
    # enough. Dropping them keeps the grid from opening with a screenful
    # of blanks.
    width = 0
    for row in merged:
        for i in range(len(row) - 1, -1, -1):
            if row[i] != "":
                width = max(width, i + 1)
                break

    hit_column_cap = width > MAX_COLUMNS_PER_SHEET
    width = min(width, MAX_COLUMNS_PER_SHEET)

    # `max_row` is what the file claims, which is the honest denominator
    # for "showing 500 of N" even though we stopped reading early.
    declared_rows = value_sheet.max_row or 0

    if width == 0 or not merged:
        return (
            SheetPreview(
                name=str(value_sheet.title),
                headers=[],
                rows=[],
                total_rows=max(0, declared_rows - 1) if declared_rows else 0,
                truncated=False,
                truncated_by=None,
            ),
            0,
        )

    hit_row_cap = len(merged) > row_limit
    body_source = merged[1 : row_limit + 1]

    headers = _normalize_headers(_fit(merged[0], width))
    rows = [_fit(row, width) for row in body_source]

    total_rows = max(len(rows), declared_rows - 1 if declared_rows else 0)

    return (
        SheetPreview(
            name=str(value_sheet.title),
            headers=headers,
            rows=rows,
            total_rows=total_rows,
            truncated=hit_row_cap or hit_column_cap,
            truncated_by="columns" if hit_column_cap else ("rows" if hit_row_cap else None),
        ),
        len(rows) * width,
    )


def _take_rows(sheet, limit: int) -> list[tuple[Any, ...]]:
    """Pull at most `limit` rows, stopping the stream rather than reading
    the whole sheet to throw most of it away."""
    out: list[tuple[Any, ...]] = []
    try:
        for row in sheet.iter_rows(values_only=True):
            out.append(row)
            if len(out) >= limit:
                break
    except Exception as e:
        # A sheet can fail mid-stream on a malformed shared-string or
        # date. Keep whatever was read — a partial preview beats none.
        logger.warning("Stopped reading sheet early: %s", e)
    return out


def _merge_row(value_row: Iterable[Any], formula_row: Iterable[Any]) -> list[str]:
    """Format one row, falling back to formula text where no value was cached."""
    values = list(value_row)
    formulas = list(formula_row)
    out: list[str] = []
    for i, value in enumerate(values):
        if value is None and i < len(formulas):
            candidate = formulas[i]
            if isinstance(candidate, str) and candidate.startswith("="):
                out.append(candidate)
                continue
        out.append(_format_cell(value))
    return out


def _fit(row: list[str], width: int) -> list[str]:
    """Pad or clip a row to the sheet's grid width."""
    if len(row) == width:
        return row
    if len(row) > width:
        return row[:width]
    return row + [""] * (width - len(row))


def _normalize_headers(row: list[str]) -> list[str]:
    """Column labels, with the spreadsheet letter standing in for a blank."""
    return [h.strip() or get_column_letter(i + 1) for i, h in enumerate(row)]


def _format_cell(value: Any) -> str:
    """Render one cell value the way it should read in the grid.

    The formatting rules a spreadsheet applies (currency, percentages,
    thousands separators) live in the cell's `number_format`, which
    `read_only` + `values_only` does not carry. Rather than half-apply
    them, values are rendered plainly and predictably. The one thing
    worth care is float noise: a cell holding 0.1 + 0.2 must not preview
    as 0.30000000000000004.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        # Checked before int: bool is a subclass of int in Python, and
        # Excel shows these as TRUE/FALSE.
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, (float, Decimal)):
        return _format_number(float(value))
    if isinstance(value, datetime):
        # openpyxl returns a datetime for every date-formatted cell, so a
        # plain date arrives with a midnight time that was never in the
        # file. Showing it would invent precision.
        if value.hour == value.minute == value.second == value.microsecond == 0:
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    return str(value)


def _format_number(value: float) -> str:
    """Format a float without exposing binary-float noise."""
    if value != value or value in (float("inf"), float("-inf")):
        return str(value)
    if value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    # 15 significant digits is the most a float carries; formatting to it
    # collapses 0.30000000000000004 to 0.3 while leaving genuine
    # precision alone.
    text = f"{value:.15g}"
    return text
