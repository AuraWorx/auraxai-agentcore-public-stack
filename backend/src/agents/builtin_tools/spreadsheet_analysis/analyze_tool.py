"""Analyze spreadsheet files using Code Interpreter.

Factory function creates a context-bound tool that downloads tabular files
from S3, pushes them to Code Interpreter, and executes Python code for analysis.
"""

import logging
import os
from typing import Any, Dict, Optional

import boto3
from strands import tool

from .list_files_tool import _get_kb_files, _get_session_files

logger = logging.getLogger(__name__)

MAX_OUTPUT_CHARS = 10000  # ~2500 tokens — safe margin under context limits


def _truncate_output(text: str) -> str:
    """Truncate tool output to prevent blowing the LLM context window."""
    if not text or len(text) <= MAX_OUTPUT_CHARS:
        return text
    return text[:MAX_OUTPUT_CHARS] + f"\n\n... (output truncated — {len(text):,} chars total, showing first {MAX_OUTPUT_CHARS:,})"


def _get_code_interpreter_id() -> Optional[str]:
    """Get Code Interpreter ID from environment or SSM."""
    ci_id = os.getenv("AGENTCORE_CODE_INTERPRETER_ID")
    if ci_id:
        return ci_id
    try:
        project_name = os.getenv("PROJECT_NAME", "strands-agent-chatbot")
        environment = os.getenv("ENVIRONMENT", "dev")
        region = os.getenv("AWS_REGION", "us-west-2")
        ssm = boto3.client("ssm", region_name=region)
        response = ssm.get_parameter(Name=f"/{project_name}/{environment}/agentcore/code-interpreter-id")
        return response["Parameter"]["Value"]
    except Exception:
        return None


def make_analyze_tool(
    assistant_id: Optional[str],
    session_id: str,
    user_id: str,
):
    """Create an analyze_spreadsheet tool bound to the given context."""

    @tool
    def analyze_spreadsheet(
        filename: str,
        python_code: str,
        output_filename: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Analyze a spreadsheet file using Python code in Code Interpreter.

        Downloads the specified file and loads it into a sandboxed Python environment
        for analysis. Use pandas, numpy, matplotlib, and seaborn.

        File loading:
        - XLSX files: converted to CSV in the sandbox. Use pd.read_csv('filename.csv')
          where filename.csv is the original name with .csv extension.
        - CSV files: available directly. Use pd.read_csv('filename.csv')

        Best for: aggregations, filtering, trends, comparisons, statistics, charts.
        For simple factual lookups, use the knowledge base search instead.

        Args:
            filename: Name of the file to analyze (from list_available_files results).
            python_code: Python code to execute. Always use pd.read_csv() to load the data.
                        The CSV filename is the original name with .csv extension
                        (e.g., "report.xlsx" becomes "report.csv").
                        Available libraries: pandas, numpy, matplotlib, seaborn, openpyxl.
            output_filename: Optional PNG filename if generating a chart.
                           Must end with .png. Example: "chart.png"

        Returns:
            Analysis results as text, and optionally a chart image.
        """
        from bedrock_agentcore.tools.code_interpreter_client import CodeInterpreter

        # 1. Validate Code Interpreter is available
        ci_id = _get_code_interpreter_id()
        if not ci_id:
            return {"content": [{"text": "❌ Code Interpreter is not configured. Contact your administrator."}], "status": "error"}

        # 2. Find the file in accessible sources
        file_info = _find_file(filename, assistant_id, session_id)
        if not file_info:
            return {"content": [{"text": f"❌ File '{filename}' not found or not accessible. Use list_available_files to see available files."}], "status": "error"}

        # 3. Download from S3
        try:
            file_bytes = _download_file(file_info)
        except Exception as e:
            return {"content": [{"text": f"❌ Failed to download file: {e}"}], "status": "error"}

        # 4. Push file to Code Interpreter
        content_type = file_info.get("content_type", "")
        is_xlsx = "spreadsheetml" in content_type or filename.lower().endswith(".xlsx")

        region = os.getenv("AWS_REGION", "us-west-2")
        code_interpreter = CodeInterpreter(region)

        try:
            code_interpreter.start(identifier=ci_id)

            if is_xlsx:
                # Push XLSX as base64, decode in sandbox
                import base64
                b64_content = base64.b64encode(file_bytes).decode("ascii")
                csv_filename = os.path.splitext(filename)[0] + ".csv"

                # Write base64 file and decode script
                code_interpreter.invoke("writeFiles", {"content": [
                    {"path": "_encoded.b64", "text": b64_content},
                ]})
                # Decode and convert to CSV inside the sandbox
                bootstrap_code = f"""
import base64, io, csv
from openpyxl import load_workbook

with open('_encoded.b64', 'r') as f:
    raw = base64.b64decode(f.read())

wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
with open('{csv_filename}', 'w', newline='') as out:
    writer = csv.writer(out)
    for row in ws.iter_rows(values_only=True):
        if all(cell is None for cell in row):
            continue
        writer.writerow([str(cell) if cell is not None else '' for cell in row])
wb.close()
print(f"Converted to {csv_filename}")
"""
                resp = code_interpreter.invoke("executeCode", {"code": bootstrap_code, "language": "python", "clearContext": False})
                # Check for bootstrap errors
                for event in resp.get("stream", []):
                    result = event.get("result", {})
                    if result.get("isError", False):
                        error_msg = result.get("structuredContent", {}).get("stderr", "Unknown error")
                        return {"content": [{"text": f"❌ Failed to convert XLSX in sandbox:\n```\n{error_msg[:500]}\n```"}], "status": "error"}
            else:
                # CSV — push directly as text
                csv_filename = filename if filename.lower().endswith(".csv") else os.path.splitext(filename)[0] + ".csv"
                try:
                    csv_text = file_bytes.decode("utf-8")
                except UnicodeDecodeError:
                    csv_text = file_bytes.decode("utf-8", errors="replace")
                code_interpreter.invoke("writeFiles", {"content": [{"path": csv_filename, "text": csv_text}]})

            # Execute user code
            response = code_interpreter.invoke("executeCode", {
                "code": python_code,
                "language": "python",
                "clearContext": False,
            })

            # Process execution result
            execution_output = ""
            for event in response.get("stream", []):
                result = event.get("result", {})
                if result.get("isError", False):
                    error_msg = result.get("structuredContent", {}).get("stderr", "Unknown error")
                    return {
                        "content": [{"text": f"❌ Code execution failed:\n```\n{error_msg[:1000]}\n```\n\nNote: The file is available as CSV at path: `{csv_filename}`. Use `pd.read_csv('{csv_filename}')` to load it."}],
                        "status": "error",
                    }
                # Extract stdout
                stdout = result.get("structuredContent", {}).get("stdout", "")
                if stdout:
                    execution_output += stdout

            # 6. Download chart if requested
            if output_filename and output_filename.endswith(".png"):
                try:
                    dl_response = code_interpreter.invoke("readFiles", {"paths": [output_filename]})
                    file_content = None
                    for event in dl_response.get("stream", []):
                        result = event.get("result", {})
                        if "content" in result:
                            for block in result["content"]:
                                if "data" in block:
                                    file_content = block["data"]
                                elif "resource" in block and "blob" in block["resource"]:
                                    file_content = block["resource"]["blob"]
                                if file_content:
                                    break
                        if file_content:
                            break

                    if file_content:
                        return {
                            "content": [
                                {"text": execution_output or f"✅ Analysis complete. Chart: {output_filename}"},
                                {"image": {"format": "png", "source": {"bytes": file_content}}},
                            ],
                            "status": "success",
                        }
                except Exception as e:
                    logger.warning(f"Failed to download chart {output_filename}: {e}")

            return {
                "content": [{"text": _truncate_output(execution_output) or "✅ Code executed successfully (no output)."}],
                "status": "success",
            }

        finally:
            try:
                code_interpreter.stop()
            except Exception:
                pass

    return analyze_spreadsheet


def _find_file(filename: str, assistant_id: Optional[str], session_id: str) -> Optional[Dict[str, Any]]:
    """Find a file by name in accessible sources. Returns file info or None."""
    if assistant_id:
        for f in _get_kb_files(assistant_id):
            if f["filename"] == filename:
                return f

    for f in _get_session_files(session_id):
        if f["filename"] == filename:
            return f

    return None


def _download_file(file_info: Dict[str, Any]) -> bytes:
    """Download file bytes from S3."""
    region = os.environ.get("AWS_REGION", "us-west-2")
    s3 = boto3.client("s3", region_name=region)

    if file_info["source"] == "knowledge_base":
        bucket = os.environ.get("S3_ASSISTANTS_DOCUMENTS_BUCKET_NAME")
        if not bucket:
            raise ValueError("S3_ASSISTANTS_DOCUMENTS_BUCKET_NAME not configured")
    else:
        bucket = file_info.get("s3_bucket")
        if not bucket:
            raise ValueError("S3 bucket not found in file metadata")

    response = s3.get_object(Bucket=bucket, Key=file_info["s3_key"])
    return response["Body"].read()
